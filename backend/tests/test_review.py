from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient

from backend.app import app

from uuid import uuid4

import backend.app as app_module
from backend.models import Extraction, Requirement, Skill

client = TestClient(app)

SYNTHETIC_TEXT = (
    "Company: Synthetic Practice Co\n"
    "Role: Junior Developer\n"
    "Required: Python\n"
    "Preferred: SQL"
)

def _extract_text():
    response = client.post(
        "/api/team/v1/jd/extract-text",
        json={
            "text": SYNTHETIC_TEXT,
            "source_file": "jd_a.txt",
        },
    )
    assert response.status_code == 200, response.text
    return response.json()


def test_extract_text_defaults_to_manual_needs_review():
    result = _extract_text()

    assert result["source_type"] == "jd"
    assert result["source_text"] == SYNTHETIC_TEXT
    assert result["review_status"] == "needs_review"
    assert result["extraction_method"] == "manual"
    assert result["skills"] == []


def test_extract_uploaded_docx_returns_manual_needs_review():
    document = Document()
    document.add_paragraph("Company: Synthetic Practice Co")
    document.add_paragraph("Role: Junior Developer")
    document.add_paragraph("Required: Python")
    content = BytesIO()
    document.save(content)

    response = client.post(
        "/api/team/v1/jd/extract",
        files={
            "file": (
                "jd_a.docx",
                content.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert response.status_code == 200, response.text
    result = response.json()
    assert result["review_status"] == "needs_review"
    assert result["extraction_method"] == "manual"
    assert "Required: Python" in result["source_text"]


def test_extract_upload_rejects_unsupported_file_type():
    response = client.post(
        "/api/team/v1/jd/extract",
        files={"file": ("notes.txt", b"synthetic text", "text/plain")},
    )

    assert response.status_code == 415
    assert "error" in response.json()


def test_review_route_keeps_correction_and_increments_revision(monkeypatch):
    _mock_reviewable_extraction(monkeypatch)
    extraction = _extract_text()
    extraction["requirement_groups"][1]["requirement_type"] = "mandatory"

    response = client.post("/api/team/v1/jd/review", json=extraction)

    assert response.status_code == 200, response.text
    result = response.json()
    assert result["review_status"] == "reviewed"
    assert result["revision"] == 2
    assert result["source_text"] == SYNTHETIC_TEXT
    assert result["requirement_groups"][1]["requirement_type"] == "mandatory"


def test_review_route_rejects_forged_quote(monkeypatch):
    _mock_reviewable_extraction(monkeypatch)
    extraction = _extract_text()
    extraction["skills"][0]["evidence"] = "Invented: Java"

    response = client.post("/api/team/v1/jd/review", json=extraction)

    assert response.status_code == 422
    assert "error" in response.json()


def test_review_route_rejects_replacement_source_text():
    extraction = _extract_text()
    extraction["source_text"] = "A different source document."

    response = client.post("/api/team/v1/jd/review", json=extraction)

    assert response.status_code == 422
    assert "error" in response.json()


def test_review_route_rejects_stale_revision():
    original = _extract_text()
    first_review = client.post("/api/team/v1/jd/review", json=original)
    assert first_review.status_code == 200, first_review.text

    stale_review = client.post("/api/team/v1/jd/review", json=original)

    assert stale_review.status_code == 409
    assert "error" in stale_review.json()

def _mock_reviewable_extraction(monkeypatch):
    def fake_extract_document(source_type, source_file, source_text, consent=False):
        extraction_id = f"test-{uuid4().hex[:12]}"
        python_id = f"{extraction_id}-skill-1"
        sql_id = f"{extraction_id}-skill-2"

        return Extraction(
            schema_version="TEAM-V1",
            id=extraction_id,
            revision=1,
            source_type=source_type,
            source_file=source_file,
            source_text=source_text,
            company="Synthetic Practice Co",
            role="Junior Developer",
            name="",
            email="",
            skills=[
                Skill(
                    skill_id=python_id,
                    skill_name="Python",
                    category_code="COD",
                    evidence="Required: Python",
                    confidence="medium",
                    source_type="document",
                    source_id=extraction_id,
                    source_page=None,
                    claim_status="claimed",
                ),
                Skill(
                    skill_id=sql_id,
                    skill_name="SQL",
                    category_code="SQL",
                    evidence="Preferred: SQL",
                    confidence="medium",
                    source_type="document",
                    source_id=extraction_id,
                    source_page=None,
                    claim_status="claimed",
                ),
            ],
            requirement_groups=[
                Requirement(
                    requirement_id=f"{extraction_id}-req-1",
                    skill_ids=[python_id],
                    operator="AND",
                    requirement_type="mandatory",
                    evidence="Required: Python",
                ),
                Requirement(
                    requirement_id=f"{extraction_id}-req-2",
                    skill_ids=[sql_id],
                    operator="AND",
                    requirement_type="preferred",
                    evidence="Preferred: SQL",
                ),
            ],
            education=[],
            projects=[],
            experience=[],
            internships=[],
            hackathons=[],
            certifications=[],
            achievements=[],
            review_status="needs_review",
            extraction_method="manual",
            warnings=[],
        )

    monkeypatch.setattr(app_module, "extract_document", fake_extract_document)
