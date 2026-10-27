from io import BytesIO
from pathlib import Path
import zipfile

from docx import Document
from pypdf import PdfReader


MAX_FILE_SIZE_BYTES = 10 * 1024 * 1024


def _read_pdf(content: bytes) -> str:
    if not content.startswith(b"%PDF-"):
        raise ValueError("The file does not have a valid PDF signature.")

    try:
        reader = PdfReader(BytesIO(content), strict=True)

        if reader.is_encrypted:
            raise ValueError("Password-protected PDFs are not supported.")

        pages = [page.extract_text() or "" for page in reader.pages]
    except ValueError:
        raise
    except Exception as error:
        raise ValueError("The PDF is corrupt or unreadable.") from error

    return "\n".join(pages)


def _read_docx(content: bytes) -> str:
    if not zipfile.is_zipfile(BytesIO(content)):
        raise ValueError("The file does not have a valid DOCX signature.")

    try:
        document = Document(BytesIO(content))
    except Exception as error:
        raise ValueError("The DOCX file is corrupt or unreadable.") from error

    parts = [
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]

    # Include text in tables as well as ordinary paragraphs.
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append("\t".join(cells))

    return "\n".join(parts)


def read_document(filename: str, content: bytes) -> str:
    """Extract selectable text from a PDF or DOCX document."""
    if not isinstance(filename, str) or not filename.strip():
        raise ValueError("A filename is required.")

    extension = Path(filename).suffix.lower()
    if extension not in {".pdf", ".docx"}:
        raise ValueError("Only PDF and DOCX files are supported.")

    if not isinstance(content, (bytes, bytearray)):
        raise ValueError("The file content must be bytes.")

    content = bytes(content)

    if len(content) > MAX_FILE_SIZE_BYTES:
        raise ValueError("The file is larger than the 10 MB limit.")

    if not content:
        raise ValueError("The file is empty.")

    if extension == ".pdf":
        text = _read_pdf(content)
    else:
        text = _read_docx(content)

    if not text.strip():
        raise ValueError("The document contains no readable text.")

    return text.strip()